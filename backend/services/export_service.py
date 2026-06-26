"""
导出服务 - 处理Excel/Word导出
"""

import io
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from typing import List, Dict, Any


class ExportService:
    """导出服务类"""
    
    @staticmethod
    def export_scores_to_excel(summary: List[Dict[str, Any]]) -> io.BytesIO:
        """
        导出成绩到Excel
        
        Args:
            summary: 成绩汇总数据
            
        Returns:
            Excel文件的BytesIO对象
        """
        df = pd.DataFrame(summary)
        
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='成绩汇总', index=False)
            
            # 调整列宽
            worksheet = writer.sheets['成绩汇总']
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 30)
                worksheet.column_dimensions[column_letter].width = adjusted_width
        
        output.seek(0)
        return output
    
    @staticmethod
    def export_student_report(username: str, profile: Dict, submissions: List) -> io.BytesIO:
        """
        导出学生个人报告到Word
        
        Args:
            username: 学生用户名
            profile: 画像数据
            submissions: 提交记录列表
            
        Returns:
            Word文件的BytesIO对象
        """
        doc = Document()
        
        # 标题
        title = doc.add_heading(f'{username} 能力画像报告', 0)
        title.alignment = 1  # 居中
        
        # 综合得分
        doc.add_heading('综合得分', level=1)
        doc.add_paragraph(f'{profile.get("overall", 0)} 分')
        
        # 各维度得分
        doc.add_heading('各维度得分', level=1)
        dimensions = profile.get("dimensions", {})
        for key, dim_data in dimensions.items():
            status = dim_data.get("status", "pending")
            if status == "evaluated":
                doc.add_paragraph(f'{dim_data.get("name", key)}: {dim_data.get("score", 0)} 分')
            else:
                doc.add_paragraph(f'{dim_data.get("name", key)}: 待评测')
        
        # 提交记录
        doc.add_heading('提交记录', level=1)
        for sub in submissions[:10]:
            p = doc.add_paragraph()
            p.add_run(f'任务: {sub.get("task_title", "未知")}\n').bold = True
            p.add_run(f'时间: {sub.get("submit_time", "未知")}\n')
            if sub.get("weighted_total"):
                p.add_run(f'得分: {sub["weighted_total"]} 分')
        
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output
    
    @staticmethod
    def export_dimension_report(dimensions_data: Dict) -> io.BytesIO:
        """
        导出维度分析报告
        """
        doc = Document()
        doc.add_heading('维度分析报告', 0)
        
        for dim_name, scores in dimensions_data.items():
            doc.add_heading(dim_name, level=1)
            doc.add_paragraph(f'平均分: {sum(scores)/len(scores):.1f}' if scores else '暂无数据')
            doc.add_paragraph(f'提交次数: {len(scores)}')
        
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output