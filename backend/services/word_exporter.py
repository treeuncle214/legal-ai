"""
Word文档导出服务 - 测评报告
"""
import os
import re
from datetime import datetime
from typing import Dict, List

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from backend.config import EXPORT_DIR


# ============================================================
# 1. 文本清洗
# ============================================================

def clean_text(text: str) -> str:
    """彻底清洗文本"""
    if not text:
        return ""
    
    # 清理 Markdown 标记
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    text = re.sub(r'\*([^*]+)\*', r'\1', text)
    text = re.sub(r'__([^_]+)__', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    
    # 移除不可见字符
    text = text.replace('\u200b', '')
    text = text.replace('\ufeff', '')
    text = text.replace('\u200c', '')
    text = text.replace('\u200d', '')
    text = text.replace('\u00a0', ' ')
    
    # 统一引号
    text = text.replace("‘", "'")
    text = text.replace("’", "'")
    text = text.replace("“", '"')
    text = text.replace("”", '"')
    text = text.replace("（", "(")
    text = text.replace("）", ")")
    
    # 清理多余空白
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = text.strip()
    
    return text


# ============================================================
# 2. 强制设置字体
# ============================================================

def set_run_font(run, font_name: str = "宋体", size: float = 10.5, bold: bool = False):
    """强制清除字体格式，重新设置"""
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    
    r = run._element
    rPr = r.get_or_add_rPr()
    
    for child in list(rPr):
        tag = child.tag
        if tag.endswith("rFonts") or tag.endswith("b") or tag.endswith("bCs"):
            rPr.remove(child)
    
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)
    
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)


def set_cell_font(cell, font_name: str = "宋体", size: float = 10.5, bold: bool = False):
    """强制清除单元格所有段落的字体格式"""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.clear()
        
        if not paragraph.runs:
            run = paragraph.add_run()
        else:
            run = paragraph.runs[0]
        
        set_run_font(run, font_name, size, bold)


# ============================================================
# 3. 段落添加函数
# ============================================================

def add_paragraph(doc, text: str, font_size: int = 10.5, bold: bool = False,
                  first_line_indent: float = 0, space_after: float = 6,
                  font_name: str = '宋体') -> None:
    if not text:
        return
    
    text = clean_text(text)
    if not text:
        return
    
    p = doc.add_paragraph()
    
    for run in p.runs:
        run.clear()
    
    run = p.add_run(text)
    set_run_font(run, font_name, font_size, bold)
    
    if first_line_indent > 0:
        p.paragraph_format.first_line_indent = Inches(first_line_indent)
    else:
        p.paragraph_format.first_line_indent = None
    
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.5


def add_heading(doc, text: str, level: int = 1) -> None:
    text = clean_text(text)
    if not text:
        return
    
    p = doc.add_paragraph()
    
    for run in p.runs:
        run.clear()
    
    run = p.add_run(text)
    font_size = 16 if level == 1 else 14
    set_run_font(run, "黑体", font_size, bold=False)
    
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)


# ============================================================
# 4. 主要导出函数
# ============================================================

def export_report_to_word(
    student_name: str,
    student_id: str,
    task_title: str,
    dimension_scores: Dict[str, float],
    indicator_scores: Dict[str, float],
    report_data: Dict,
    student_college: str = "",
    student_major: str = "",
    total_score: float = 0.0
) -> str:
    """
    导出测评报告为Word文档
    """
    os.makedirs(EXPORT_DIR, exist_ok=True)
    
    doc = Document()
    
    # 设置默认样式
    try:
        style = doc.styles['Normal']
        style.font.name = '宋体'
        style.font.size = Pt(10.5)
        rPr = style._element.rPr
        if rPr is not None:
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn('w:eastAsia'), '宋体')
            rFonts.set(qn('w:ascii'), '宋体')
            rFonts.set(qn('w:hAnsi'), '宋体')
    except:
        pass
    
    # ========== 标题 ==========
    p = doc.add_paragraph()
    for run in p.runs:
        run.clear()
    
    run = p.add_run(f'《{clean_text(task_title)}》测评意见')
    set_run_font(run, "黑体", 22, bold=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # ========== 学生信息（含学院、专业） ==========
    student_info = f'学生：{clean_text(student_name)}（{student_id}）'
    if student_college:
        student_info += f'  学院：{clean_text(student_college)}'
    if student_major:
        student_info += f'  专业：{clean_text(student_major)}'
    add_paragraph(doc, student_info, font_size=12, bold=False)
    add_paragraph(doc, f'日期：{datetime.now().strftime("%Y年%m月%d日")}', font_size=12, bold=False)
    
    doc.add_paragraph()
    
    # ========== 一、总体评价 ==========
    add_heading(doc, '一、总体评价', level=1)
    
    overall = report_data.get("overall_evaluation", "暂无总体评价内容。")
    add_paragraph(doc, overall, font_size=10.5, bold=False, first_line_indent=0.3, space_after=12)
    
    doc.add_paragraph()
    
    # ========== 二、得分情况 ==========
    add_heading(doc, '二、得分情况', level=1)

    dimension_names = {
        "ai_retrieval": "AI融合智能检索能力",
        "critical": "批判性评估能力",
        "ethics": "伦理合规辨识能力",
        "integration": "信息整合应用能力"
    }

    # 2.1 维度得分表格
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'

    headers = ['维度', '得分', '等级']
    for i, header in enumerate(headers):
        set_cell_font(table.rows[0].cells[i], "宋体", 10.5, bold=True)
        table.rows[0].cells[i].text = header

    # 遍历所有维度并显示
    for dim_key, dim_name in dimension_names.items():
        score = dimension_scores.get(dim_key, 0)
        level = "优" if score >= 85 else "良" if score >= 75 else "合格" if score >= 55 else "不合格"
        
        row_cells = table.add_row().cells
        set_cell_font(row_cells[0], "宋体", 10.5, bold=False)
        row_cells[0].text = dim_name
        set_cell_font(row_cells[1], "宋体", 10.5, bold=False)
        row_cells[1].text = f"{score:.2f}"
        set_cell_font(row_cells[2], "宋体", 10.5, bold=False)
        row_cells[2].text = level

    # 总分行
    avg_level = "优" if total_score >= 85 else "良" if total_score >= 75 else "合格" if total_score >= 55 else "不合格"
    row_cells = table.add_row().cells
    set_cell_font(row_cells[0], "宋体", 10.5, bold=True)
    row_cells[0].text = "总分"
    set_cell_font(row_cells[1], "宋体", 10.5, bold=True)
    row_cells[1].text = f"{total_score:.2f}"
    set_cell_font(row_cells[2], "宋体", 10.5, bold=True)
    row_cells[2].text = avg_level

    doc.add_paragraph()
    
    # 2.2 二级指标得分（按维度分组，更清晰）
    indicator_names = {
        "A1": "检索目标拆解",
        "A2": "检索策略设计",
        "A3": "AI工具融合应用",
        "A4": "检索策略优化",
        "B1": "信息来源评估",
        "B2": "AI内容验证",
        "B3": "争议与分歧分析",
        "C1": "风险类型识别",
        "C2": "价值综合判断",
        "C3": "风险处理方式",
        "D1": "信息分类与组织",
        "D2": "综合分析与决策",
        "D3": "局限反思"
    }
    
    # 获取指标满分
    indicator_max_scores = {}
    if report_data.get("indicator_max_scores"):
        indicator_max_scores = report_data["indicator_max_scores"]
    else:
        for key in indicator_names.keys():
            indicator_max_scores[key] = 10
    
    dim_indicator_map = {
        "ai_retrieval": ['A1', 'A2', 'A3', 'A4'],
        "critical": ['B1', 'B2', 'B3'],
        "ethics": ['C1', 'C2', 'C3'],
        "integration": ['D1', 'D2', 'D3']
    }
    
    # 添加二级指标标题
    add_paragraph(doc, "二级指标得分：", font_size=10.5, bold=False)
    
    # 按维度分别创建小表格（更清晰）
    for dim_key, dim_name in dimension_names.items():
        dim_indicators = dim_indicator_map.get(dim_key, [])
        if not dim_indicators:
            continue
        
        # 维度标题行（在表格外）
        add_paragraph(doc, f"【{dim_name}】", font_size=10.5, bold=True, space_after=4)
        
        # 创建该维度的指标表格（4列）
        ind_table = doc.add_table(rows=1, cols=4)
        ind_table.style = 'Table Grid'
        
        # 表头
        ind_headers = ['指标', '指标名称', '得分', '满分']
        for i, header in enumerate(ind_headers):
            set_cell_font(ind_table.rows[0].cells[i], "宋体", 10.5, bold=True)
            ind_table.rows[0].cells[i].text = header
        
        for ind_key in dim_indicators:
            score = indicator_scores.get(ind_key, 0)
            max_score = indicator_max_scores.get(ind_key, 10)
            name = indicator_names.get(ind_key, ind_key)
            
            row_cells = ind_table.add_row().cells
            set_cell_font(row_cells[0], "宋体", 10.5, bold=False)
            row_cells[0].text = ind_key
            set_cell_font(row_cells[1], "宋体", 10.5, bold=False)
            row_cells[1].text = name
            set_cell_font(row_cells[2], "宋体", 10.5, bold=False)
            
            # ✅ 关键修改：如果分数为0，得分和满分都显示"本次未涉及"
            if score > 0:
                row_cells[2].text = f"{score:.1f} / {max_score}分"
                set_cell_font(row_cells[3], "宋体", 10.5, bold=False)
                row_cells[3].text = str(max_score)
            else:
                row_cells[2].text = "本次未涉及"
                set_cell_font(row_cells[3], "宋体", 10.5, bold=False)
                row_cells[3].text = "本次未涉及"
        
        doc.add_paragraph()
    
    doc.add_paragraph()
    
    # ========== 三、问题反馈与学习建议 ==========
    add_heading(doc, '三、问题反馈与学习建议', level=1)
    
    issue_feedback = report_data.get("issue_feedback", [])
    
    if not issue_feedback:
        add_paragraph(doc, "暂无问题反馈，建议教师根据具体情况添加。", font_size=10.5, bold=False)
    else:
        for issue in issue_feedback:
            title = issue.get("title", "")
            performance = issue.get("performance", "")
            improvement = issue.get("improvement", "")
            
            add_paragraph(doc, title, font_size=12, bold=False, space_after=6)
            
            if performance:
                add_paragraph(doc, performance, font_size=10.5, bold=False, first_line_indent=0.3, space_after=6)
            
            if improvement:
                add_paragraph(doc, improvement, font_size=10.5, bold=False, first_line_indent=0.3, space_after=12)
            
            doc.add_paragraph()
    
    # ========== 保存文件 ==========
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"测评报告_{student_id}_{timestamp}.docx"
    file_path = os.path.join(EXPORT_DIR, filename)
    doc.save(file_path)
    
    return file_path