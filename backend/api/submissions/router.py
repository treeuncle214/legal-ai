"""
提交评分 API 主路由
"""
import os
import uuid
import asyncio
import io
import logging
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Response
from sqlalchemy.orm import Session
from datetime import datetime
from docx import Document

from backend.api.deps import get_db, get_current_user, get_current_student, get_student_class_id, get_teacher_class_ids
from backend.database import add_submission, update_scores, get_task
from backend.database.submissions import can_submit, get_submission_count
from backend.core.scorer import score_submission
from backend.schemas.submission import TextSubmissionRequest
from backend.schemas.common import Response as APIResponse
from backend.config import UPLOAD_DIR

# 导入子模块路由
from backend.api.submissions.trigger import router as trigger_router
from backend.api.submissions.batch import router as batch_router
from backend.api.submissions.queries import router as queries_router
from backend.api.submissions.utils import get_download_response
from backend.api.submissions.task_wrapper import TaskWrapper

logger = logging.getLogger(__name__)

MAX_FILE_SIZE = 10 * 1024 * 1024
router = APIRouter(prefix="/api", tags=["提交评分"])

# 包含子路由
router.include_router(trigger_router)
router.include_router(batch_router)
router.include_router(queries_router)


# ========== OPTIONS 处理 ==========
@router.options("/submissions/text")
async def options_submissions_text():
    return Response(status_code=200, headers={
        "Access-Control-Allow-Origin": "*",
        "Access-Control-Allow-Methods": "POST, OPTIONS",
        "Access-Control-Allow-Headers": "Content-Type, Authorization",
    })

@router.options("/submissions/word")
async def options_submissions_word():
    return Response(status_code=200, headers={
        "Access-Control-Allow-Origin": "*",
        "Access-Control-Allow-Methods": "POST, OPTIONS",
        "Access-Control-Allow-Headers": "Content-Type, Authorization",
    })


# ========== 提交接口 ==========

@router.post("/submissions/text", response_model=APIResponse)
async def submit_text(
    submission_data: TextSubmissionRequest,
    current_user: dict = Depends(get_current_student),
    db: Session = Depends(get_db)
):
    try:
        task_dict = get_task(submission_data.task_id)
        if not task_dict:
            raise HTTPException(status_code=404, detail="任务不存在")
        
        student_class_id = get_student_class_id(current_user)
        if task_dict.get("class_id") != student_class_id:
            raise HTTPException(status_code=403, detail="无权提交此任务")
        
        task_obj = TaskWrapper(task_dict)
        can_sub, check_message = can_submit(current_user["username"], submission_data.task_id, task_obj)
        if not can_sub:
            raise HTTPException(status_code=400, detail=check_message)
        
        content = submission_data.final_output
        score_result = score_submission(task_dict, {"final_output": content, "submit_type": "text"})
        submission_id = add_submission(
            task_id=submission_data.task_id,
            student_username=current_user["username"],
            process_log=submission_data.process_log,
            ai_interaction_log=submission_data.ai_interaction_log,
            final_output=content,
            tools_used=",".join(submission_data.tools_used) if submission_data.tools_used else "",
            submit_type="text",
            ai_score_status="completed",
            ai_scored=True,
            ai_scored_at=datetime.now().isoformat(),
            ai_scored_by="system"
        )
        scores_dict = {
            "score_ai_retrieval": score_result["dimension_scores"].get("ai_retrieval", 0),
            "score_critical": score_result["dimension_scores"].get("critical", 0),
            "score_ethics": score_result["dimension_scores"].get("ethics", 0),
            "score_integration": score_result["dimension_scores"].get("integration", 0),
            "ai_comment": score_result.get("comment", "AI评分完成"),
            "ai_score_detail": str(score_result.get("indicator_grades", {})) if "indicator_grades" in score_result else str(score_result.get("module_scores", {}))
        }
        update_scores(submission_id, scores_dict)
        submission_count = get_submission_count(current_user["username"], submission_data.task_id)
        remaining_after = max(0, task_obj.max_submissions - submission_count)
        return APIResponse(
            data={
                "submission_id": submission_id,
                "status": "completed",
                "dimension_scores": score_result["dimension_scores"],
                "ai_comment": score_result.get("comment", ""),
                "remaining_submissions": remaining_after,
                "max_submissions": task_obj.max_submissions,
                "submitted_count": submission_count
            },
            message=f"提交成功！剩余提交次数：{remaining_after}"
        )
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/submissions/word", response_model=APIResponse)
async def submit_word(
    task_id: int = Form(...),
    file1: UploadFile = File(...),
    file2: UploadFile = File(...),
    current_user: dict = Depends(get_current_student),
    db: Session = Depends(get_db)
):
    """
    上传两个Word文档：
    - file1: AI交互记录 → ai_interaction_log
    - file2: 作业正文 → process_log + final_output
    """
    for f in [file1, file2]:
        f.file.seek(0, 2)
        if f.file.tell() > MAX_FILE_SIZE:
            raise HTTPException(status_code=400, detail=f"文件 {f.filename} 超过大小限制")
        f.file.seek(0)
        if not f.filename.lower().endswith('.docx'):
            raise HTTPException(status_code=400, detail=f"{f.filename} 不是 .docx 格式")

    os.makedirs(UPLOAD_DIR, exist_ok=True)

    try:
        content1 = await file1.read()
        content2 = await file2.read()
        
        doc1 = Document(io.BytesIO(content1))
        doc2 = Document(io.BytesIO(content2))
        
        ai_interaction_text = "\n".join([p.text for p in doc1.paragraphs]).strip()
        doc2_text = "\n".join([p.text for p in doc2.paragraphs]).strip()
        
        if not ai_interaction_text or not doc2_text:
            raise HTTPException(status_code=400, detail="两个文档内容均不能为空")
        
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"解析文档失败: {str(e)}")

    task = get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    
    student_class_id = get_student_class_id(current_user)
    if task.get("class_id") != student_class_id:
        raise HTTPException(status_code=403, detail="无权提交此任务")
    
    task_obj = TaskWrapper(task)
    can_sub, check_message = can_submit(current_user["username"], task_id, task_obj)
    if not can_sub:
        raise HTTPException(status_code=400, detail=check_message)

    original_name1 = file1.filename
    original_name2 = file2.filename
    
    unique1 = f"{uuid.uuid4().hex}.docx"
    unique2 = f"{uuid.uuid4().hex}.docx"
    path1 = os.path.join(UPLOAD_DIR, unique1)
    path2 = os.path.join(UPLOAD_DIR, unique2)
    with open(path1, "wb") as f:
        f.write(content1)
    with open(path2, "wb") as f:
        f.write(content2)

    combined = f"【AI交互记录】\n{ai_interaction_text}\n\n【作业正文】\n{doc2_text}"

    submission_id = add_submission(
        task_id=task_id,
        student_username=current_user["username"],
        process_log=doc2_text,
        ai_interaction_log=ai_interaction_text,
        final_output=doc2_text,
        submit_type="word",
        word_file_path=f"{unique1},{unique2}",
        word_content=combined,
        ai_score_status="pending",
        ai_scored=False,
        original_filenames=f"{original_name1},{original_name2}"
    )

    # ❌ 移除自动评分：不再自动触发AI评分

    return APIResponse(
        data={"submission_id": submission_id, "message": "提交成功，等待教师AI评分"},
        message="提交成功！等待教师进行AI评分"
    )


# ========== 文件下载 ==========
@router.get("/download/{filename}")
async def download_file(
    filename: str,
    current_user: dict = Depends(get_current_user)
):
    return get_download_response(filename, current_user)