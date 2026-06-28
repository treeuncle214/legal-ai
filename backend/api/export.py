# backend/api/export.py
"""
数据导出 API
"""

from fastapi import APIRouter, Depends, HTTPException, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import io
import pandas as pd
from docx import Document

from backend.api.deps import get_db, get_current_teacher, get_teacher_class_ids
from backend.database import get_all_submissions_summary, get_submission
from backend.database.tasks import get_task

router = APIRouter(prefix="/api", tags=["数据导出"])


@router.get("/export/scores")
async def export_scores(
    current_user: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db)
):
    """导出全班成绩为Excel（仅当前教师班级）"""
    teacher_class_ids = get_teacher_class_ids(current_user)
    if not teacher_class_ids:
        raise HTTPException(status_code=400, detail="您没有管理的班级")
    
    # 获取所有提交摘要，然后按班级过滤
    summary = get_all_submissions_summary()
    
    # 过滤出属于教师班级的学生
    from backend.database.engine import get_db_connection
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        # 获取教师班级的所有学生用户名
        cursor.execute("""
            SELECT u.username 
            FROM users u
            JOIN user_class uc ON u.id = uc.user_id
            WHERE uc.class_id IN ({})
            AND u.role = 'student'
        """.format(','.join('?' * len(teacher_class_ids))), teacher_class_ids)
        student_usernames = [row[0] for row in cursor.fetchall()]
    finally:
        conn.close()
    
    # 过滤
    filtered_summary = [s for s in summary if s["student_username"] in student_usernames]
    
    # 转换为DataFrame
    df = pd.DataFrame(filtered_summary)
    
    # 写入Excel
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='成绩汇总', index=False)
    
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=scores.xlsx"}
    )


@router.get("/export/student_report/{username}")
async def export_student_report(
    username: str,
    current_user: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db)
):
    """导出学生个人报告为Word（仅当前教师班级）"""
    # 验证该学生在教师班级
    from backend.database.engine import get_db_connection
    teacher_class_ids = get_teacher_class_ids(current_user)
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT uc.class_id 
            FROM user_class uc
            JOIN users u ON u.id = uc.user_id
            WHERE u.username = ?
        """, (username,))
        student_class_rows = cursor.fetchall()
        student_class_ids = [row[0] for row in student_class_rows]
        if not any(cid in teacher_class_ids for cid in student_class_ids):
            raise HTTPException(status_code=403, detail="无权导出该学生报告")
    finally:
        conn.close()
    
    from backend.database import calculate_profile, get_submissions_by_student
    from backend.config import get_dimension_name
    
    profile = calculate_profile(username)
    submissions = get_submissions_by_student(username)
    
    # 创建Word文档
    doc = Document()
    doc.add_heading(f'{username} 能力画像报告', 0)
    
    doc.add_heading('综合得分', level=1)
    doc.add_paragraph(f'{profile["overall"]} 分')
    
    doc.add_heading('各维度得分', level=1)
    for key, dim_data in profile["dimensions"].items():
        if dim_data["status"] == "evaluated":
            doc.add_paragraph(f'{dim_data["name"]}: {dim_data["score"]} 分')
        else:
            doc.add_paragraph(f'{dim_data["name"]}: 待评测')
    
    doc.add_heading('提交记录', level=1)
    for sub in submissions[:10]:
        doc.add_paragraph(f'任务: {sub.get("task_title", "未知")}, 时间: {sub["submit_time"]}')
    
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={username}_report.docx"}
    )