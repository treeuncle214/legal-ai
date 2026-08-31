"""
Excel Sheet 构建函数
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from backend.config import SCORING_DIMENSIONS
from backend.api.scores.helpers import get_level


def create_sheet1(ws, students, tasks, all_scores, student_names, task_info):
    """Sheet 1: 任务总览"""
    # ✅ 添加学院和专业列
    headers = ["学号", "姓名", "学院", "专业"]
    task_ids = []
    for task in tasks:
        headers.append(f"{task.title}")
        task_ids.append(task.id)
    headers.append("平均分")
    headers.append("等级")
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    row_num = 2
    student_scores = {}
    
    for student in students:
        username = student["username"]
        display_name = student["display_name"]
        college = student.get("college", "")
        major = student.get("major", "")
        student_scores[username] = []
        
        ws.cell(row=row_num, column=1, value=username)
        ws.cell(row=row_num, column=2, value=display_name)
        ws.cell(row=row_num, column=3, value=college or "-")
        ws.cell(row=row_num, column=4, value=major or "-")
        
        col = 5  # 从第5列开始是任务成绩
        total_sum = 0
        valid_count = 0
        
        for task in tasks:
            score_data = all_scores.get(username, {}).get(task.id)
            if score_data:
                score = score_data["total"]
                ws.cell(row=row_num, column=col, value=score)
                total_sum += score
                valid_count += 1
                student_scores[username].append(score)
            else:
                ws.cell(row=row_num, column=col, value="-")
            col += 1
        
        avg_score = round(total_sum / valid_count, 2) if valid_count > 0 else 0
        level = get_level(avg_score) if valid_count > 0 else "未提交"
        
        ws.cell(row=row_num, column=col, value=avg_score)
        ws.cell(row=row_num, column=col + 1, value=level)
        
        row_num += 1
    
    if row_num > 2:
        ws.cell(row=row_num, column=1, value="📊 全班平均")
        ws.cell(row=row_num, column=2, value="")
        ws.cell(row=row_num, column=3, value="")
        ws.cell(row=row_num, column=4, value="")
        
        col = 5
        for task in tasks:
            scores = []
            for student in students:
                score_data = all_scores.get(student["username"], {}).get(task.id)
                if score_data:
                    scores.append(score_data["total"])
            avg = round(sum(scores) / len(scores), 2) if scores else 0
            ws.cell(row=row_num, column=col, value=avg)
            col += 1
        
        all_avg = []
        for student in students:
            student_avg = student_scores.get(student["username"], [])
            if student_avg:
                all_avg.append(sum(student_avg) / len(student_avg))
        total_avg = round(sum(all_avg) / len(all_avg), 2) if all_avg else 0
        ws.cell(row=row_num, column=col, value=total_avg)
        
        summary_font = Font(bold=True)
        summary_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        for c in range(1, col + 2):
            cell = ws.cell(row=row_num, column=c)
            cell.font = summary_font
            cell.fill = summary_fill
    
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 16  # 学院列
    ws.column_dimensions["D"].width = 16  # 专业列


def create_sheet4(ws, tasks, all_scores, student_names, task_info):
    """Sheet 2: 各任务汇总（显示全班各维度百分制平均分）"""
    dimension_keys = [d["key"] for d in SCORING_DIMENSIONS]
    dimension_names = {d["key"]: d["name"] for d in SCORING_DIMENSIONS}
    dimension_short_names = {
        "ai_retrieval": "AI融合检索",
        "critical": "批判性评估",
        "ethics": "伦理合规",
        "integration": "信息整合"
    }
    
    headers = ["任务名称", "任务类型"]
    for dim_key in dimension_keys:
        headers.append(f"{dimension_short_names[dim_key]}(平均)")
    headers.append("全班平均分")
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    row_num = 2
    for task in tasks:
        ws.cell(row=row_num, column=1, value=task.title)
        ws.cell(row=row_num, column=2, value=task_info[task.id]["type"])
        
        col = 3
        all_total_scores = []
        
        for dim_key in dimension_keys:
            dim_scores = []
            for student in student_names.keys():
                score_data = all_scores.get(student, {}).get(task.id)
                if score_data:
                    dim_score = score_data["dimension_scores"].get(dim_key, 0)
                    # ✅ 只统计有分数的学生（>0表示该维度被考核）
                    if dim_score > 0:
                        dim_scores.append(dim_score)
            
            # ✅ 百分制平均分
            avg = round(sum(dim_scores) / len(dim_scores), 2) if dim_scores else 0
            ws.cell(row=row_num, column=col, value=avg)
            col += 1
        
        # ✅ 全班总分平均（每个学生的作业总分是百分制）
        for student in student_names.keys():
            score_data = all_scores.get(student, {}).get(task.id)
            if score_data:
                all_total_scores.append(score_data["total"])
        
        task_avg = round(sum(all_total_scores) / len(all_total_scores), 2) if all_total_scores else 0
        ws.cell(row=row_num, column=col, value=task_avg)
        
        row_num += 1
    
    # 汇总行
    if row_num > 2:
        ws.cell(row=row_num, column=1, value="📊 全部任务汇总")
        ws.cell(row=row_num, column=2, value="")
        
        col = 3
        for dim_key in dimension_keys:
            dim_scores = []
            for task in tasks:
                task_row = tasks.index(task) + 2
                val = ws.cell(row=task_row, column=col).value
                if val and isinstance(val, (int, float)) and val > 0:
                    dim_scores.append(val)
            avg = round(sum(dim_scores) / len(dim_scores), 2) if dim_scores else 0
            ws.cell(row=row_num, column=col, value=avg)
            col += 1
        
        all_avgs = []
        for task in tasks:
            task_row = tasks.index(task) + 2
            val = ws.cell(row=task_row, column=col).value
            if val and isinstance(val, (int, float)) and val > 0:
                all_avgs.append(val)
        total_avg = round(sum(all_avgs) / len(all_avgs), 2) if all_avgs else 0
        ws.cell(row=row_num, column=col, value=total_avg)
        
        summary_font = Font(bold=True)
        summary_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        for c in range(1, col + 1):
            cell = ws.cell(row=row_num, column=c)
            cell.font = summary_font
            cell.fill = summary_fill
    
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 14

def get_dimension_max_score(task_id: int, dim_key: str, all_scores: dict) -> int:
    """
    获取任务某维度的满分
    从 all_scores 中的 max_scores 计算该维度的满分之和
    """
    # 获取维度对应的指标
    dim_indicators = {
        "ai_retrieval": ["A1", "A2", "A3", "A4"],
        "critical": ["B1", "B2", "B3"],
        "ethics": ["C1", "C2", "C3"],
        "integration": ["D1", "D2", "D3"]
    }
    
    indicators = dim_indicators.get(dim_key, [])
    total_max = 0
    
    # 从任一学生的成绩中获取 max_scores
    for student_scores in all_scores.values():
        task_scores = student_scores.get(task_id)
        if task_scores and task_scores.get("max_scores"):
            for ind_key in indicators:
                total_max += task_scores["max_scores"].get(ind_key, 10)
            return total_max
    
    # 默认每个指标10分
    return len(indicators) * 10