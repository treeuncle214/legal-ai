"""
Excel 导出主逻辑
"""

from openpyxl import Workbook
from fastapi import HTTPException
from sqlalchemy.orm import Session

from backend.api.scores.sheets import create_sheet1, create_sheet4
from backend.api.scores.helpers import get_student_names, get_task_info


def export_class_scores_to_excel(
    db: Session,
    class_id: int,
    class_name: str,
    students: list,
    tasks: list,
    all_scores: dict
) -> bytes:
    """
    导出全班成绩为 Excel 文件
    只包含两个 Sheet：
    - Sheet 1: 任务总览
    - Sheet 4: 各任务汇总
    """
    try:
        wb = Workbook()
        
        # ✅ 删除默认的 Sheet
        default_sheet = wb.active
        wb.remove(default_sheet)
        
        # ✅ 获取辅助数据
        student_names = get_student_names(students)
        task_info = get_task_info(tasks)
        
        # ✅ Sheet 1: 任务总览
        ws1 = wb.create_sheet("任务总览")
        create_sheet1(ws1, students, tasks, all_scores, student_names, task_info)
        
        # ✅ Sheet 4: 各任务汇总
        ws4 = wb.create_sheet("各任务汇总")
        create_sheet4(ws4, tasks, all_scores, student_names, task_info)
        
        # 保存到字节流
        import io
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出失败: {str(e)}")