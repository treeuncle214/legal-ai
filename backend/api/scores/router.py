"""
成绩总览 API 路由定义
"""

from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
import logging
import io
from datetime import datetime
from urllib.parse import quote

from backend.api.deps import get_db
from backend.database.models import User, Task, Submission, SubmissionScore
from backend.database.classes import get_class_students, get_class
from backend.api.deps import get_current_user, get_teacher_class_ids
from backend.config import SCORING_DIMENSIONS
from backend.api.scores.analytics import get_class_analytics_data
from backend.api.scores.helpers import get_all_indicator_max_scores
from backend.api.scores.sheets import create_sheet1, create_sheet4

try:
    from openpyxl import Workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

router = APIRouter(prefix="/api/scores", tags=["成绩总览"])
logger = logging.getLogger(__name__)


@router.get("/class-analytics/{class_id}")
async def get_class_analytics(
    class_id: int,
    task_type: Optional[str] = Query(None, description="课堂练习/任务实践/综合考察"),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    """
    获取班级学情分析数据
    """
    class_obj = get_class(class_id)
    if not class_obj:
        raise HTTPException(status_code=404, detail="班级不存在")
    
    if current_user["role"] != "admin":
        teacher_class_ids = get_teacher_class_ids(current_user)
        if class_id not in teacher_class_ids:
            raise HTTPException(status_code=403, detail="无权查看此班级数据")
    
    data = get_class_analytics_data(class_id, task_type, db, current_user)
    return data


@router.get("/export-class-scores/{class_id}")
async def export_class_scores(
    class_id: int,
    task_type: Optional[str] = Query(None, description="课堂练习/任务实践/综合考察"),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    """
    导出全班成绩 Excel（只包含两个 Sheet：任务总览、各任务汇总）
    """
    if not OPENPYXL_AVAILABLE:
        raise HTTPException(status_code=500, detail="openpyxl 未安装，请运行 pip install openpyxl")
    
    class_info = get_class(class_id)
    if not class_info:
        raise HTTPException(status_code=404, detail="班级不存在")
    
    class_name = class_info.get("name", f"班级{class_id}")
    
    # 权限检查
    if current_user["role"] != "admin":
        teacher_class_ids = get_teacher_class_ids(current_user)
        if class_id not in teacher_class_ids:
            raise HTTPException(
                status_code=403,
                detail=f"无权导出此班级数据。您的班级: {teacher_class_ids}，请求班级: {class_id}"
            )
    
    students = get_class_students(class_id)
    
    # 为每个学生补充学院和专业信息
    for student in students:
        user = db.query(User).filter(User.username == student["username"]).first()
        if user:
            student["college"] = user.college or ""
            student["major"] = user.major or ""
        else:
            student["college"] = ""
            student["major"] = ""
    
    student_usernames = [s["username"] for s in students]
    student_names = {s["username"]: s["display_name"] for s in students}
    
    tasks = db.query(Task).filter(
        Task.class_id == class_id,
        Task.is_active == 1
    )
    if task_type:
        tasks = tasks.filter(Task.task_type == task_type)
    tasks = tasks.order_by(Task.created_at).all()
    
    if not tasks:
        raise HTTPException(status_code=404, detail="该班级暂无任务")
    
    all_scores = {}
    task_info = {}
    
    for task in tasks:
        # ✅ 关键修复：获取每个学生该任务的最新提交（而不是所有提交）
        subquery = db.query(
            Submission.student_username,
            func.max(Submission.id).label("max_id")
        ).filter(
            Submission.task_id == task.id,
            Submission.student_username.in_(student_usernames),
            Submission.score_published == 1,
            Submission.is_reviewed == 1
        ).group_by(Submission.student_username).subquery()
        
        # 获取最新提交记录
        submissions = db.query(Submission).join(
            subquery,
            Submission.id == subquery.c.max_id
        ).all()
        
        task_info[task.id] = {
            "title": task.title,
            "type": task.task_type or "任务实践",
            "created_at": task.created_at
        }
        
        for sub in submissions:
            if sub.student_username not in all_scores:
                all_scores[sub.student_username] = {}
            
            indicator_scores = db.query(SubmissionScore).filter(
                SubmissionScore.submission_id == sub.id
            ).all()
            indicator_dict = {s.indicator_key: s.score for s in indicator_scores}
            max_scores = get_all_indicator_max_scores(task.id, db)
            
            # ✅ 关键修复：优先使用 final_score（教师审批后的百分制得分）
            # 否则使用 score_*（AI原始百分制得分）
            dimension_scores = {}
            for dim in SCORING_DIMENSIONS:
                key = dim["key"]
                final_score = getattr(sub, f"final_score_{key}", None)
                ai_score = getattr(sub, f"score_{key}", 0)
                
                if final_score is not None and float(final_score) > 0:
                    dimension_scores[key] = round(float(final_score), 2)
                else:
                    dimension_scores[key] = round(float(ai_score), 2)
            
            # ✅ 关键修复：总分直接使用数据库中的 total_score（百分制）
            total = sub.total_score if sub.total_score else 0
            
            all_scores[sub.student_username][task.id] = {
                "total": round(float(total), 2),
                "dimension_scores": dimension_scores,
                "indicator_scores": indicator_dict,
                "max_scores": max_scores
            }
    
    # 创建工作簿 - 只创建两个 Sheet
    wb = Workbook()
    wb.remove(wb.active)
    
    # Sheet 1: 任务总览
    ws1 = wb.create_sheet("任务总览", 0)
    create_sheet1(ws1, students, tasks, all_scores, student_names, task_info)
    
    # Sheet 2: 各任务汇总
    ws4 = wb.create_sheet("各任务汇总", 1)
    create_sheet4(ws4, tasks, all_scores, student_names, task_info)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"全班成绩_{class_name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    encoded_filename = quote(filename)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
        }
    )