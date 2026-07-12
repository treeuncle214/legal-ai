# backend/api/submissions.py
"""
提交和评分 API
"""

import os
import tempfile
import shutil
import uuid
import asyncio
import io
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Response
from sqlalchemy.orm import Session
from typing import List
from datetime import datetime
from fastapi.responses import FileResponse
import logging
from docx import Document

from backend.api.deps import get_db, get_current_user, get_current_student, get_student_class_id, get_teacher_class_ids
from backend.database import (
    add_submission, update_scores, get_task, get_submissions_by_student
)
from backend.database.submissions import (
    can_submit, get_submission_count, get_submission, update_ai_score_status,
    delete_submission
)
from backend.database.tasks import get_task as get_task_db
from backend.core.scorer import score_submission
from backend.schemas.submission import TextSubmissionRequest
from backend.schemas.common import Response as APIResponse
from backend.config import UPLOAD_DIR
from backend.database.models import SubmissionScore
from backend.database.engine import SessionLocal



logger = logging.getLogger(__name__)

MAX_FILE_SIZE = 10 * 1024 * 1024
router = APIRouter(prefix="/api", tags=["提交评分"])


class TaskWrapper:
    def __init__(self, data: dict):
        self.id = data.get('id')
        self.title = data.get('title')
        self.description = data.get('description')
        self.due_date = data.get('due_date')
        self.task_type = data.get('task_type', '任务实践')
        self.enabled_indicators = data.get('enabled_indicators', '')
        self.max_submissions = data.get('max_submissions', 3)
        self.allow_after_deadline = data.get('allow_after_deadline', 0)
        self.class_id = data.get('class_id')  # 新增
    
    def get_enabled_indicators_list(self) -> List[str]:
        if not self.enabled_indicators:
            return []
        return [s.strip() for s in self.enabled_indicators.split(',') if s.strip()]


# ========== OPTIONS 处理 ==========
@router.options("/submissions/text")
async def options_submissions_text():
    return Response(status_code=200, headers={
        "Access-Control-Allow-Origin": "*",
        "Access-Control-Allow-Methods": "POST, OPTIONS",
        "Access-Control-Allow-Headers": "Content-Type, Authorization",
    })

@router.options("/submissions/word")
async def options_submissions_word():
    return Response(status_code=200, headers={
        "Access-Control-Allow-Origin": "*",
        "Access-Control-Allow-Methods": "POST, OPTIONS",
        "Access-Control-Allow-Headers": "Content-Type, Authorization",
    })


# ========== AI 评分异步任务 ==========
# backend/api/submissions.py
# 修改 perform_scoring 函数

async def perform_scoring(submission_id: int, task_dict: dict, content: str, submit_type: str):
    try:
        logger.info(f"开始为提交 {submission_id} 进行AI评分")
        update_ai_score_status(submission_id, "scoring")
        score_result = score_submission(task_dict, {"final_output": content, "submit_type": submit_type})
        
        # 准备评分数据
        scores_dict = {
            "score_ai_retrieval": score_result["dimension_scores"].get("ai_retrieval", 0),
            "score_critical": score_result["dimension_scores"].get("critical", 0),
            "score_ethics": score_result["dimension_scores"].get("ethics", 0),
            "score_integration": score_result["dimension_scores"].get("integration", 0),
            "ai_comment": score_result.get("comment", "AI评分完成"),
            "ai_score_status": "completed",
            "ai_score_detail": score_result.get("indicator_grades", {}) if "indicator_grades" in score_result else score_result.get("module_scores", {})
        }
        update_scores(submission_id, scores_dict)
        
        # ========== 新增：存储每个指标的评分详情 ==========
        try:
            from backend.database.models import SubmissionScore
            from backend.database.engine import SessionLocal
            
            db = SessionLocal()
            try:
                # 获取指标得分（从 score_result 中提取）
                indicator_scores = score_result.get("indicator_scores", {})
                indicator_levels = score_result.get("indicator_levels", {})
                indicator_comments = score_result.get("indicator_comments", {})
                
                for key, score in indicator_scores.items():
                    indicator_score = SubmissionScore(
                        submission_id=submission_id,
                        indicator_key=key,
                        score=score,
                        level=indicator_levels.get(key, "合格"),
                        comment=indicator_comments.get(key, "")
                    )
                    db.add(indicator_score)
                db.commit()
                logger.info(f"提交 {submission_id} 指标评分详情已保存")
            finally:
                db.close()
        except Exception as e:
            logger.error(f"保存指标评分详情失败: {e}")
        
        logger.info(f"提交 {submission_id} AI评分完成")
    except Exception as e:
        logger.error(f"提交 {submission_id} AI评分失败: {str(e)}", exc_info=True)
        update_ai_score_status(submission_id, "failed", error_message=str(e))
        update_scores(submission_id, {"ai_comment": "AI评分失败，请教师手动批改", "ai_score_status": "failed"})



# ========== 文本提交（保留，但前端已不再使用） ==========
@router.post("/submissions/text", response_model=APIResponse)
async def submit_text(
    submission_data: TextSubmissionRequest,
    current_user: dict = Depends(get_current_student),
    db: Session = Depends(get_db)
):
    try:
        task_dict = get_task(submission_data.task_id)
        if not task_dict:
            raise HTTPException(status_code=404, detail="任务不存在")
        
        # 验证任务属于学生所在班级
        student_class_id = get_student_class_id(current_user)
        if task_dict.get("class_id") != student_class_id:
            raise HTTPException(status_code=403, detail="无权提交此任务")
        
        task_obj = TaskWrapper(task_dict)
        can_sub, check_message = can_submit(current_user["username"], submission_data.task_id, task_obj)
        if not can_sub:
            raise HTTPException(status_code=400, detail=check_message)
        
        content = submission_data.final_output
        score_result = score_submission(task_dict, {"final_output": content, "submit_type": "text"})
        submission_id = add_submission(
            task_id=submission_data.task_id,
            student_username=current_user["username"],
            process_log=submission_data.process_log,
            ai_interaction_log=submission_data.ai_interaction_log,
            final_output=content,
            tools_used=",".join(submission_data.tools_used) if submission_data.tools_used else "",
            submit_type="text",
            ai_score_status="completed"
        )
        scores_dict = {
            "score_ai_retrieval": score_result["dimension_scores"].get("ai_retrieval", 0),
            "score_critical": score_result["dimension_scores"].get("critical", 0),
            "score_ethics": score_result["dimension_scores"].get("ethics", 0),
            "score_integration": score_result["dimension_scores"].get("integration", 0),
            "ai_comment": score_result.get("comment", "AI评分完成"),
            "ai_score_detail": str(score_result.get("indicator_grades", {})) if "indicator_grades" in score_result else str(score_result.get("module_scores", {}))
        }
        update_scores(submission_id, scores_dict)
        submission_count = get_submission_count(current_user["username"], submission_data.task_id)
        remaining_after = max(0, task_obj.max_submissions - submission_count)
        return APIResponse(
            data={
                "submission_id": submission_id,
                "status": "completed",
                "dimension_scores": score_result["dimension_scores"],
                "ai_comment": score_result.get("comment", ""),
                "remaining_submissions": remaining_after,
                "max_submissions": task_obj.max_submissions,
                "submitted_count": submission_count
            },
            message=f"提交成功！剩余提交次数：{remaining_after}"
        )
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


# ========== Word 双文件提交 ==========
@router.post("/submissions/word", response_model=APIResponse)
async def submit_word(
    task_id: int = Form(...),
    file1: UploadFile = File(...),
    file2: UploadFile = File(...),
    current_user: dict = Depends(get_current_student),
    db: Session = Depends(get_db)
):
    """上传两个Word文档：file1=检索过程，file2=最终结果"""
    # 校验文件
    for f in [file1, file2]:
        f.file.seek(0, 2)
        if f.file.tell() > MAX_FILE_SIZE:
            raise HTTPException(status_code=400, detail=f"文件 {f.filename} 超过大小限制")
        f.file.seek(0)
        if not f.filename.lower().endswith('.docx'):
            raise HTTPException(status_code=400, detail=f"{f.filename} 不是 .docx 格式")

    os.makedirs(UPLOAD_DIR, exist_ok=True)

    # 读取并提取文本
    try:
        content1 = await file1.read()
        content2 = await file2.read()
        doc1 = Document(io.BytesIO(content1))
        doc2 = Document(io.BytesIO(content2))
        process_log_text = "\n".join([p.text for p in doc1.paragraphs]).strip()
        final_output_text = "\n".join([p.text for p in doc2.paragraphs]).strip()
        if not process_log_text or not final_output_text:
            raise HTTPException(status_code=400, detail="两个文档内容均不能为空")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"解析文档失败: {str(e)}")

    # 获取任务
    task_dict = get_task(task_id)
    if not task_dict:
        raise HTTPException(status_code=404, detail="任务不存在")
    
    # 验证任务属于学生所在班级
    student_class_id = get_student_class_id(current_user)
    if task_dict.get("class_id") != student_class_id:
        raise HTTPException(status_code=403, detail="无权提交此任务")
    
    task_obj = TaskWrapper(task_dict)

    # 检查提交次数
    can_sub, check_message = can_submit(current_user["username"], task_id, task_obj)
    if not can_sub:
        raise HTTPException(status_code=400, detail=check_message)

    # 保存文件
    original_name1 = file1.filename
    original_name2 = file2.filename
    
    unique1 = f"{uuid.uuid4().hex}.docx"
    unique2 = f"{uuid.uuid4().hex}.docx"
    path1 = os.path.join(UPLOAD_DIR, unique1)
    path2 = os.path.join(UPLOAD_DIR, unique2)
    with open(path1, "wb") as f:
        f.write(content1)
    with open(path2, "wb") as f:
        f.write(content2)

    # 合并内容给AI
    combined = f"【检索过程记录】\n{process_log_text}\n\n【最终检索结果】\n{final_output_text}"

    # 保存原始文件名（逗号分隔，与存储文件顺序一致）
    original_filenames = f"{original_name1},{original_name2}"

    # 插入数据库，状态 pending
    submission_id = add_submission(
        task_id=task_id,
        student_username=current_user["username"],
        process_log=process_log_text,
        final_output=final_output_text,
        submit_type="word",
        word_file_path=f"{unique1},{unique2}",
        word_content=combined,
        ai_score_status="pending",
        original_filenames=original_filenames  # 新增：保存原始文件名
    )

    # 异步执行AI评分
    asyncio.create_task(
        perform_scoring(
            submission_id=submission_id,
            task_dict=task_dict,
            content=combined,
            submit_type="word"
        )
    )

    # 返回成功（不含评分）
    return APIResponse(
        data={"submission_id": submission_id, "message": "提交成功，等待教师批改"},
        message="提交成功！请等待教师批改后查看成绩"
    )

# ========== 其余查询接口 ==========
@router.get("/submissions/{submission_id}/score")
async def get_submission_score(
    submission_id: int,
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    submission = get_submission(submission_id)
    if not submission:
        raise HTTPException(status_code=404, detail="提交记录不存在")
    if current_user["role"] != "teacher" and submission["student_username"] != current_user["username"]:
        raise HTTPException(status_code=403, detail="无权查看此提交")
    score_status = submission.get("ai_score_status", "pending")
    if score_status == "completed":
        return {
            "status": "completed",
            "data": {
                "dimension_scores": {
                    "ai_retrieval": submission.get("score_ai_retrieval", 0),
                    "critical": submission.get("score_critical", 0),
                    "ethics": submission.get("score_ethics", 0),
                    "integration": submission.get("score_integration", 0)
                },
                "ai_comment": submission.get("ai_comment", ""),
                "ai_score_detail": submission.get("ai_score_detail", {})
            }
        }
    elif score_status == "scoring":
        return {"status": "scoring", "message": "AI评分进行中，请稍后"}
    elif score_status == "failed":
        return {"status": "failed", "message": f"AI评分失败: {submission.get('ai_score_error', '未知错误')}"}
    else:
        return {"status": "pending", "message": "等待评分开始"}


@router.get("/submissions/student/{username}", response_model=APIResponse[list])
async def get_student_submissions(
    username: str,
    current_user: dict = Depends(get_current_student),
    db: Session = Depends(get_db)
):
    if current_user["username"] != username:
        raise HTTPException(status_code=403, detail="只能查看自己的提交记录")
    from backend.database.submissions import get_student_submissions_without_scores
    submissions = get_student_submissions_without_scores(username)
    return APIResponse(data=submissions)


@router.get("/submissions/task/{task_id}", response_model=APIResponse[list])
async def get_task_submissions(
    task_id: int,
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    from backend.database import get_submissions_by_task_all
    from backend.database.tasks import get_task as get_task_db
    
    task = get_task_db(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    
    if current_user["role"] == "teacher":
        # 教师：验证任务属于自己班级
        teacher_class_ids = get_teacher_class_ids(current_user)
        if task.get("class_id") not in teacher_class_ids:
            raise HTTPException(status_code=403, detail="无权查看此任务")
        submissions = get_submissions_by_task_all(task_id)
    else:
        # 学生：只查看自己的提交，且任务属于自己班级
        student_class_id = get_student_class_id(current_user)
        if task.get("class_id") != student_class_id:
            raise HTTPException(status_code=403, detail="无权查看此任务")
        submissions = get_submissions_by_student(current_user["username"])
        submissions = [s for s in submissions if s["task_id"] == task_id]
    
    return APIResponse(data=submissions)


@router.get("/submissions/published", response_model=APIResponse[list])
async def get_published_scores(
    current_user: dict = Depends(get_current_student),
    db: Session = Depends(get_db)
):
    from backend.database.submissions import get_published_submissions_for_student
    submissions = get_published_submissions_for_student(current_user["username"])
    result = []
    for sub in submissions:
        result.append({
            "id": sub["id"],
            "task_title": sub.get("task_title", "未知任务"),
            "weighted_total": sub.get("weighted_total", 0),
            "teacher_comment": sub.get("teacher_comment", ""),
            "submit_time": sub.get("submit_time"),
            "is_reviewed": sub.get("is_reviewed", 1),
        })
    return APIResponse(data=result)


@router.get("/submissions/{submission_id}", response_model=APIResponse)
async def get_submission_detail(
    submission_id: int,
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    submission = get_submission(submission_id)
    if not submission:
        raise HTTPException(status_code=404, detail="提交记录不存在")
    
    # 权限检查
    if current_user["role"] != "teacher" and submission["student_username"] != current_user["username"]:
        raise HTTPException(status_code=403, detail="无权查看此提交")
    
    # 教师查看时验证提交属于自己班级
    if current_user["role"] == "teacher":
        task = get_task_db(submission.get("task_id"))
        teacher_class_ids = get_teacher_class_ids(current_user)
        if task and task.get("class_id") not in teacher_class_ids:
            raise HTTPException(status_code=403, detail="无权查看此提交")
    
    # ✅ 新增：查询指标级评分
    from backend.database.models import SubmissionScore
    from backend.database.engine import SessionLocal
    from backend.config import SCORING_DIMENSIONS
    
    # 获取所有13个指标名称映射
    indicator_names = {}
    for dim in SCORING_DIMENSIONS:
        for ind in dim.get("sub_indicators", []):
            indicator_names[ind["key"]] = ind["name"]
    
    # 查询该提交的指标评分
    db_local = SessionLocal()
    try:
        scores = db_local.query(SubmissionScore).filter(
            SubmissionScore.submission_id == submission_id
        ).all()
        
        indicator_scores = {}
        indicator_levels = {}
        indicator_comments = {}
        
        for s in scores:
            indicator_scores[s.indicator_key] = s.score
            indicator_levels[s.indicator_key] = s.level
            indicator_comments[s.indicator_key] = s.comment
        
        # 添加到返回数据
        submission["indicator_scores"] = indicator_scores
        submission["indicator_levels"] = indicator_levels
        submission["indicator_comments"] = indicator_comments
        submission["indicator_names"] = indicator_names
        
    finally:
        db_local.close()
    
    return APIResponse(data=submission)

@router.get("/submissions/remaining/{task_id}")
async def get_remaining_submissions(
    task_id: int,
    current_user: dict = Depends(get_current_student),
    db: Session = Depends(get_db)
):
    task_dict = get_task(task_id)
    if not task_dict:
        raise HTTPException(status_code=404, detail="任务不存在")
    
    # 验证任务属于学生所在班级
    student_class_id = get_student_class_id(current_user)
    if task_dict.get("class_id") != student_class_id:
        raise HTTPException(status_code=403, detail="无权查看此任务")
    
    task_obj = TaskWrapper(task_dict)
    from backend.database.submissions import get_submission_count
    submission_count = get_submission_count(current_user["username"], task_id)
    remaining = max(0, task_obj.max_submissions - submission_count)
    is_deadline_passed = False
    if task_obj.due_date and not task_obj.allow_after_deadline:
        try:
            due_date = datetime.fromisoformat(task_obj.due_date) if isinstance(task_obj.due_date, str) else task_obj.due_date
            is_deadline_passed = datetime.now() > due_date
        except:
            pass
    return {
        "remaining": remaining,
        "max_submissions": task_obj.max_submissions,
        "submitted_count": submission_count,
        "is_deadline_passed": is_deadline_passed,
        "due_date": task_obj.due_date,
        "allow_after_deadline": task_obj.allow_after_deadline == 1
    }


@router.get("/download/{filename}")
async def download_file(
    filename: str,
    current_user: dict = Depends(get_current_user)
):
    if '..' in filename or filename.startswith('/'):
        raise HTTPException(status_code=400, detail="无效的文件名")
    
    file_path = os.path.join(UPLOAD_DIR, filename)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail=f"文件不存在: {filename}")
    
    print(f"📥 下载请求: {filename}")
    
    # 查询数据库获取原始文件名
    from backend.database.submissions import get_submission_by_file_path
    submission = get_submission_by_file_path(filename)
    print(f"📊 查询结果: {submission}")
    
    original_filename = filename
    
    if submission and submission.get("original_filenames"):
        stored_files = submission["word_file_path"].split(',')
        original_files = submission["original_filenames"].split(',')
        print(f"📂 stored_files: {stored_files}")
        print(f"📂 original_files: {original_files}")
        
        if filename in stored_files:
            idx = stored_files.index(filename)
            if idx < len(original_files):
                original_filename = original_files[idx].strip()
                print(f"✅ 使用原始文件名: {original_filename}")
        else:
            print(f"⚠️ 文件名 {filename} 不在 stored_files 中")
    else:
        print(f"⚠️ 未找到 original_filenames")
    
    return FileResponse(
        file_path,
        filename=original_filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )